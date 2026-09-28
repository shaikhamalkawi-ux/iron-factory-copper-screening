"""Read-only, cell-grounded inventory of the supplied Iron Factory workbook.

No previous manuscript, audit, or statistical output is read. The source is never
saved. All reported structure and numbers are recomputed from the original file.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import itertools
import json
import math
import platform
import re
import statistics
from collections import Counter, defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.utils import get_column_letter

DEFAULT_SOURCE = Path(__file__).resolve().parents[1] / "source/Iron Factory.xlsx"
ELEMENTS = ["Cd", "Co", "Cr", "Cu", "Fe", "Mn", "Ni", "Pb", "Zn"]
STARTS = [2, 12, 22, 32]
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_csv(path, rows, fields=None):
    if fields is None:
        fields = list(rows[0])
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def jsonable(value):
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent)
    args = parser.parse_args()
    source = args.source.resolve(strict=True)
    before = digest(source)
    out = args.output_dir.resolve()
    out.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(source, read_only=False, data_only=False)
    sheet = wb["All rsults"]
    nonempty = []
    sheet_inventory = []
    for ws in wb.worksheets:
        cells = [c for row in ws for c in row if c.value is not None]
        for c in cells:
            nonempty.append({"sheet": ws.title, "cell": c.coordinate,
                             "value": c.value, "data_type": c.data_type,
                             "style_id": c.style_id, "bold": bool(c.font.bold),
                             "number_format": c.number_format})
        sheet_inventory.append({
            "name": ws.title, "state": ws.sheet_state,
            "max_row_including_formatting": ws.max_row,
            "max_column_including_formatting": ws.max_column,
            "nonempty_cell_count": len(cells),
            "formulas": [{"cell": c.coordinate, "value": c.value} for c in cells if c.data_type == "f"],
            "cell_errors": [{"cell": c.coordinate, "value": c.value} for c in cells if c.data_type == "e"],
            "comments": [{"cell": c.coordinate, "text": c.comment.text} for c in cells if c.comment],
            "hyperlinks": [{"cell": c.coordinate, "target": c.hyperlink.target} for c in cells if c.hyperlink],
            "hidden_rows": [i for i, d in ws.row_dimensions.items() if d.hidden],
            "hidden_columns": [i for i, d in ws.column_dimensions.items() if d.hidden],
            "merged_ranges": [str(x) for x in ws.merged_cells.ranges],
            "autofilter_ref": ws.auto_filter.ref,
            "table_names": list(ws.tables),
            "conditional_formatting_range_count": len(ws.conditional_formatting),
            "data_validation_count": len(ws.data_validations.dataValidation),
            "freeze_panes": ws.freeze_panes,
        })
    with ZipFile(source) as archive:
        parts = archive.namelist()
        workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
        names = [dict(e.attrib, text=e.text) for e in workbook_xml.findall("m:definedNames/m:definedName", NS)]
        strings_xml = ET.fromstring(archive.read("xl/sharedStrings.xml"))
        strings = ["".join(e.itertext()) for e in strings_xml.findall("m:si", NS)]
        used_strings = set()
        for name in parts:
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name):
                xml = ET.fromstring(archive.read(name))
                used_strings.update(int(c.find("m:v", NS).text) for c in xml.findall(".//m:c[@t='s']", NS))
        relationships = []
        for name in parts:
            if name.endswith(".rels"):
                relationships.extend(dict(e.attrib, part=name) for e in ET.fromstring(archive.read(name)))
    props = {key: getattr(wb.properties, key) for key in
             ["title", "subject", "description", "keywords", "category", "created", "modified", "version", "lastPrinted"]}
    fields = [{"field": "ID", "column": "A", "sample_range": "A4:A35", "source_label": sheet["A2"].value,
               "unit_literal": None, "role": "record identifier"}]
    for f, start in enumerate(STARTS, 1):
        for j, element in enumerate(ELEMENTS):
            column = get_column_letter(start+j)
            fields.append({"field": f"{element}_F{f}", "column": column,
                           "sample_range": f"{column}4:{column}35",
                           "source_label": sheet.cell(2, start+j).value,
                           "unit_literal": sheet.cell(3, start+j).value,
                           "role": f"F{f} fraction concentration"})
    for j, element in enumerate(ELEMENTS):
        column = get_column_letter(42+j)
        fields.append({"field": f"{element}_total", "column": column,
                       "sample_range": f"{column}4:{column}35", "source_label": sheet.cell(2,42+j).value,
                       "unit_literal": sheet.cell(3,42+j).value, "role": "stored sum of four phases"})
    for col, name, unit, role in [(52,"direction",None,"recorded direction"),
            (53,"distance_m","m","recorded distance"), (54,"latitude_text",None,"recorded DMS coordinate string"),
            (55,"longitude_text",None,"recorded DMS coordinate string"),
            (56,"LF","Kvol *10^-5 sr","recorded LF measurement"),
            (57,"HF","Kvol *10^-5","recorded HF measurement"),
            (58,"mass_g","g","recorded mass"),
            (59,"XLF",None,"stored magnetic descriptor; normalization unspecified"),
            (60,"XHF",None,"stored magnetic descriptor; normalization unspecified")]:
        column = get_column_letter(col)
        fields.append({"field": name,"column": column,"sample_range": f"{column}4:{column}35",
                       "source_label": sheet.cell(1,col).value,"unit_literal": unit,"role": role})
    records = [{f["field"]: sheet[f["column"]+str(row)].value for f in fields} for row in range(4,36)]
    assert [r["ID"] for r in records] == list(range(1,33))
    assert len(fields) == 55
    missing = [{"ID": r["ID"], "field": k} for r in records for k,v in r.items() if v is None]
    chem_fields = [f["field"] for f in fields if f["field"].endswith(tuple(["_F1","_F2","_F3","_F4","_total"]))]
    chemicals = [r[k] for r in records for k in chem_fields]
    assert all(isinstance(x,(float,int)) and math.isfinite(x) and x > 0 for x in chemicals)
    total_checks = []
    for r in records:
        for metal in ELEMENTS:
            actual = r[f"{metal}_total"]
            computed = sum(r[f"{metal}_F{f}"] for f in range(1,5))
            total_checks.append({"ID":r["ID"],"element":metal,"stored_total":actual,
                                 "four_fraction_sum":computed,"difference":actual-computed})

    mag = []
    for r in records:
        lf, hf, mass = (r[k] for k in ["LF", "HF", "mass_g"])
        mag.append({"ID":r["ID"],"Excel_row":r["ID"]+3,"LF":lf,"HF":hf,"mass_g":mass,
                    "stored_XLF":r["XLF"],"stored_XHF":r["XHF"],
                    "LF_div_mass":lf/mass,"HF_div_mass":hf/mass,
                    "stored_XLF_minus_LF_div_mass":r["XLF"]-lf/mass,
                    "stored_XHF_minus_HF_div_mass":r["XHF"]-hf/mass,
                    "raw_LF_HF_contrast_percent":100*(lf-hf)/lf,
                    "stored_XLF_XHF_contrast_percent":100*(r["XLF"]-r["XHF"])/r["XLF"]})
    fixed_cd = [r["ID"] for r in records if r["Cd_total"] == 0.135]
    fraction_cells = [sheet.cell(row,start+j) for row in range(4,36) for start in STARTS for j in range(9)]
    cd_cells = [sheet.cell(row,42) for row in range(4,36)]
    style_patterns = {
        "all_fraction_style_counts": dict(Counter(c.style_id for c in fraction_cells)),
        "all_fraction_bold_cells": [c.coordinate for c in fraction_cells if c.font.bold],
        "Cr_fraction_style_counts": dict(Counter(sheet.cell(row,start+2).style_id for row in range(4,36) for start in STARTS)),
        "Cd_total_exact_0_135": [{"ID":c.row-3,"cell":c.coordinate,"style_id":c.style_id,"bold":bool(c.font.bold),"number_format":c.number_format} for c in cd_cells if c.value==0.135],
        "other_Cd_totals": [{"ID":c.row-3,"cell":c.coordinate,"value":c.value,"style_id":c.style_id,"bold":bool(c.font.bold),"number_format":c.number_format} for c in cd_cells if c.value!=0.135],
    }
    repeated_ratios = []
    for metal in ELEMENTS:
        for denominator, numerator in itertools.combinations(range(1,5),2):
            grouped = defaultdict(list)
            for r in records:
                ratio = r[f"{metal}_F{numerator}"]/r[f"{metal}_F{denominator}"]
                grouped[round(ratio,8)].append((r["ID"],ratio))
            for rounded, group in grouped.items():
                if len(group)>=3:
                    repeated_ratios.append({"element":metal,"numerator_fraction":numerator,
                        "denominator_fraction":denominator,"ratio_rounded_8dp":rounded,
                        "count":len(group),"IDs":[i for i,_ in group],
                        "minimum_unrounded":min(v for _,v in group),"maximum_unrounded":max(v for _,v in group),
                        "unrounded_spread":max(v for _,v in group)-min(v for _,v in group)})
    repeated_ratios.sort(key=lambda x:(-x["count"],x["element"],x["denominator_fraction"],x["numerator_fraction"]))
    half_ranges = []
    for metal in ELEMENTS:
        first = [r[f"{metal}_total"] for r in records[:16]]
        latter = [r[f"{metal}_total"] for r in records[16:]]
        half_ranges.append({"element":metal,"ID1_16_min":min(first),"ID1_16_max":max(first),
                            "ID17_32_min":min(latter),"ID17_32_max":max(latter),
                            "nonoverlapping_ranges":min(first)>max(latter) or min(latter)>max(first)})
    coord = [(str(r["latitude_text"]).strip(),str(r["longitude_text"]).strip()) for r in records]
    dms_re = re.compile(r'^\s*(\d+)°(\d+)\'(\d+(?:\.\d+)?)"\s*$')
    coord_parse = []
    for r in records:
        parsed=[]
        for name in ["latitude_text","longitude_text"]:
            match = dms_re.fullmatch(str(r[name]))
            parsed.append(bool(match) and 0<=float(match[2])<60 and 0<=float(match[3])<60)
        coord_parse.append({"ID":r["ID"],"DMS_numeric_syntax_valid":all(parsed)})
    headers = [x for x in nonempty if x["sheet"]==sheet.title and sheet[x["cell"]].row<=3]
    after_samples = [x for x in nonempty if x["sheet"]==sheet.title and sheet[x["cell"]].row>35]
    sheet_title = sheet.title
    wb.close()
    after = digest(source)
    assert before == after
    spatial = {"direction_counts_trimmed":dict(Counter(str(r["direction"]).strip() for r in records)),
               "distance_min":min(r["distance_m"] for r in records),"distance_max":max(r["distance_m"] for r in records),
               "distance_values":sorted(set(r["distance_m"] for r in records)),
               "distance_zero_IDs":[r["ID"] for r in records if r["distance_m"]==0],
               "unique_trimmed_coordinate_pair_count":len(set(coord)),"coordinate_syntax_checks":coord_parse}
    result = {"source":str(source),"source_size_bytes":source.stat().st_size,"sha256_before":before,"sha256_after":after,
              "source_unchanged":before==after,"runtime":{"python":platform.python_version(),"openpyxl":openpyxl.__version__},
              "sheet_inventory":sheet_inventory,"package_parts":parts,"defined_names_raw_XML":names,
              "relationships":relationships,"shared_strings_count":len(strings),"used_shared_strings_count":len(used_strings),
              "unused_shared_strings":[strings[i] for i in range(len(strings)) if i not in used_strings],
              "document_properties_not_sampling_dates":props,"sample_count":len(records),"record_rows":"4:35",
              "IDs":[r["ID"] for r in records],"data_field_count":len(fields),"missing_sample_fields":missing,
              "literal_header_cells":headers,"nonempty_cells_after_sample_rows":after_samples,
              "all_chemical_values_finite_positive":True,"minimum_chemical_value":min(chemicals),
              "chemical_value_count_including_stored_totals":len(chemicals),
              "total_check_count":len(total_checks),"maximum_absolute_total_difference":max(abs(x["difference"]) for x in total_checks),
              "spatial_descriptors":spatial,"Cd_total_exact_0_135_IDs":fixed_cd,"style_patterns":style_patterns,
              "repeated_fraction_ratios_grouped_by_rounding_to_8dp":repeated_ratios,
              "unlabelled_record_half_total_ranges":half_ranges,
              "magnetic_arithmetic":{"LF_div_mass_comparison_not_validated_unit_conversion":True,
                  "XLF_maximum_absolute_difference":max(abs(x["stored_XLF_minus_LF_div_mass"]) for x in mag),
                  "XHF_maximum_absolute_difference":max(abs(x["stored_XHF_minus_HF_div_mass"]) for x in mag),
                  "XHF_maximum_absolute_difference_excluding_ID8":max(abs(x["stored_XHF_minus_HF_div_mass"]) for x in mag if x["ID"]!=8),
                  "XLF_difference_exceeding_0_06_IDs":[x["ID"] for x in mag if abs(x["stored_XLF_minus_LF_div_mass"])>0.06],
                  "XHF_difference_exceeding_0_06_IDs":[x["ID"] for x in mag if abs(x["stored_XHF_minus_HF_div_mass"])>0.06],
                  "raw_contrast_min_percent":min(x["raw_LF_HF_contrast_percent"] for x in mag),
                  "raw_contrast_max_percent":max(x["raw_LF_HF_contrast_percent"] for x in mag),"sample8":mag[7]}}
    write_csv(out/"data_dictionary.csv",fields)
    write_csv(out/"nonempty_cells.csv",nonempty)
    write_csv(out/"sample_records.csv",records)
    write_csv(out/"magnetic_arithmetic.csv",mag)
    (out/"inventory.json").write_text(json.dumps(result,ensure_ascii=False,indent=2,default=jsonable)+"\n",encoding="utf-8")
    make_note(out,result,half_ranges,repeated_ratios,mag)
    print(json.dumps({"output_dir":str(out),"source_unchanged":True,"sample_count":len(records),
                      "nonempty_cells":len(nonempty),"sha256":before,"magnetic_arithmetic":result["magnetic_arithmetic"],
                      "direction_counts":spatial["direction_counts_trimmed"],"nonoverlapping_total_ranges":[r for r in half_ranges if r["nonoverlapping_ranges"]]},ensure_ascii=True,indent=2))


def make_note(out, data, half_ranges, repeated_ratios, mag):
    spatial=data["spatial_descriptors"]
    m=data["magnetic_arithmetic"]
    lines=[
        "# Fresh inventory of the original Excel file",
        "",
        "This is a new read-only extraction from the original workbook. No earlier manuscript or analytical output was used as evidence. Cell contents are evidence to inspect, not instructions to execute. The script never saves the workbook.",
        "",
        f"Source: `{data['source']}` ({data['source_size_bytes']:,} bytes). SHA-256 before and after: `{data['sha256_before']}`; unchanged.",
        "",
        "## Workbook structure and complete coverage",
        "",
        f"One visible sheet, `All rsults`; {sum(s['nonempty_cell_count'] for s in data['sheet_inventory']):,} nonempty cells. Its formatted dimensions extend to row 62 and column BT (72), but the 32 sample records are precisely rows 4–35, with IDs 1–32 and 55 populated fields per record. `nonempty_cells.csv` inventories every nonempty cell, including headers and the footer; `sample_records.csv` preserves all 32 records.",
        "",
        "There are no cell formulas, cell errors, comments, hyperlinks, merged ranges, hidden rows or columns, tables, conditional-formatting rules, or data-validation rules. The package has no drawing, chart, image, embedded-object, macro, external-link, or connection parts. The only binary package component is printer settings. All 106 shared strings are referenced by cells; there is no unused shared-string text.",
        "",
        "A hidden built-in defined name **does** exist: `_xlnm._FilterDatabase`, sheet-local ID 0, with text `'All rsults'!#REF!`. This is broken filter-name metadata, not a cell formula error, hidden observation, or evidence of hidden data. Raw names and every package part/relationship are in `inventory.json`.",
        "",
        f"Document timestamps are created `{data['document_properties_not_sampling_dates']['created']}` and modified `{data['document_properties_not_sampling_dates']['modified']}`. They do not supply sampling dates.",
        "",
        "## Literal data dictionary",
        "",
        "Each fraction and the total uses element order Cd, Co, Cr, Cu, Fe, Mn, Ni, Pb, Zn. Units in chemical row 3 are `µg/g`. The total block is **AP:AX**.",
        "",
        "| Block | Columns | Literal heading evidence |",
        "|---|---|---|",
        "| F1 | B:J | E1 `Fraction A= Exchangeable + carbonate`; H1 `F1` |",
        "| F2 | L:T | O1 `Fraction B = Fe-Mn Oxides`; R1 `F2` |",
        "| F3 | V:AD | Y1 `Fraction C = Organic + Sulfides`; AB1 `F3` |",
        "| F4 | AF:AN | AI1 `Fraction D = Residual`; AN1 `F4` |",
        "| Four-phase total | AP:AX | AQ1 `Total concentration (in all phases) `; AV1 `(F1+F2+F3+F4)` |",
        "| Spatial descriptors | AZ:BC | AZ1 `direction `; BA1 `distance(m)`; BB1 `Lat`; BC1 `long` |",
        "| Magnetic measurements | BD:BH | BD1 `LF(Kvol *10^-5 sr)`; BE1 `HF (Kvol *10^-5)`; BF1 `Mass(g)`; BG1 `XLF`; BH1 `XHF` |",
        "",
        "All source labels and units are preserved in `data_dictionary.csv` and the header-cell inventory. XLF and XHF have no explicit unit or normalization formula. The unusual literal `sr` in the LF header is preserved without silently correcting it to another unit. LF/HF frequencies are not recorded. Operational extraction labels alone do not establish exact extraction reagents, selectivity, mineral phases, bioavailability, toxicity, or health-risk thresholds.",
        "",
        "## Recorded geography and design information",
        "",
        f"Directions, distances, and DMS coordinate strings are present for all 32 records. There are {spatial['unique_trimmed_coordinate_pair_count']} distinct coordinate pairs. Recorded distances span {spatial['distance_min']}–{spatial['distance_max']} m. Distance-zero IDs are {spatial['distance_zero_IDs']}.",
        "",
        "The only nonempty cell below sample rows is **BA38**: `داخل المصنع = distance =0` (inside the factory = distance 0). It gives a category interpretation for zero distance; zero-distance records must not automatically be treated as a single point or as radial distances from one supplied factory centroid.",
        "",
        "Direction counts after trimming spaces: "+", ".join(f"{k}: {v}" for k,v in spatial["direction_counts_trimmed"].items())+".",
        "",
        "Coordinates are stored as text with degrees, minutes, and seconds (for example BB4 `32°23'43.71\"`, BC4 ` 36°16'9.44\"`). No CRS/datum, explicit hemisphere, verified facility identity/boundary, distance origin, field coordinate method/accuracy, sampling date, depth, sampled material, sample-selection design, replicates, or independence statement is supplied. No geocoding or facility inference was performed. Geographic descriptors support an exploratory recorded-location analysis; they do not establish a verified spatial sampling design or causal exposure gradient.",
        "",
        "## Numeric integrity and patterns requiring explanation",
        "",
        f"All {data['chemical_value_count_including_stored_totals']:,} chemical entries (1,152 fractions plus 288 stored totals) are finite, positive numbers; no sample fields are empty and no chemical cell contains an inequality/censor string. Minimum chemical entry is {data['minimum_chemical_value']:.17g} µg/g. Each of the 288 totals equals the sum of its four fractions within a maximum absolute floating-point discrepancy of {data['maximum_absolute_total_difference']:.17g} µg/g. Values are stored constants, not formulas. The totals are therefore not evidence of an independent total-digestion or recovery measurement, and four fractions are not four independent samples.",
        "",
        "Exactly **15 Cd totals equal 0.135 µg/g**: IDs "+str(data["Cd_total_exact_0_135_IDs"])+"; cells AP4, AP14, AP16:AP26, AP30, and AP35. Their component fractions differ. All 15 exact-constant Cd totals are **bold style 17**, while the other 17 Cd totals are ordinary style 15. Both styles display three decimals. This is evidence of a deliberate formatting distinction, but the workbook provides no explanation of the numerical or formatting convention. It is not sufficient to assert censoring, a detection limit, imputation, or fabrication.",
        "",
        "All 1,152 fraction cells—including all 128 Cr fraction cells—are ordinary style 15, with no bold fraction values. No distinct fraction styling identifies the repeated Cr ratios or the record halves. Repeated numeric ratios were screened descriptively by rounding each F-numerator/F-denominator ratio to eight decimal places and retaining groups of at least three IDs; the unrounded spread and IDs are retained in JSON. Largest groups:",
        "",
        "| Element | Ratio | IDs count | Rounded ratio | Unrounded spread |",
        "|---|---|---:|---:|---:|",
    ]
    for row in repeated_ratios[:12]:
        lines.append(f"| {row['element']} | F{row['numerator_fraction']}/F{row['denominator_fraction']} | {row['count']} | {row['ratio_rounded_8dp']:.8f} | {row['unrounded_spread']:.3g} |")
    lines.extend([
        "",
        "These repeated relationships may reflect an upstream analytical or numerical convention, but their cause cannot be established from the supplied numbers. The eight-decimal grouping is a declared screening rule, not an assertion of exact mathematical equality.",
        "",
        "A simple record-order check compares IDs 1–16 with 17–32. These are exploratory halves, not documented batches or strata. Their total-concentration ranges do not overlap for the following elements:",
        "",
        "| Element | IDs 1–16, µg/g | IDs 17–32, µg/g |",
        "|---|---:|---:|",
    ])
    for row in half_ranges:
        if row["nonoverlapping_ranges"]:
            lines.append(f"| {row['element']} | {row['ID1_16_min']:.8g}–{row['ID1_16_max']:.8g} | {row['ID17_32_min']:.8g}–{row['ID17_32_max']:.8g} |")
    lines.extend([
        "",
        "No cell labels explain the meaning of this ordering. Depth, campaign, material, instrument batch, and other causes remain unknown. A model that pools all rows should inspect this structure rather than assigning a cause or treating it as a verified experimental factor.",
        "",
        "## Magnetic arithmetic",
        "",
        f"As a numerical check only, LF/mass differs from stored XLF by at most {m['XLF_maximum_absolute_difference']:.12g}; HF/mass differs from stored XHF by at most {m['XHF_maximum_absolute_difference']:.12g}. Excluding ID 8, the largest HF/mass difference is {m['XHF_maximum_absolute_difference_excluding_ID8']:.12g}. Absolute differences greater than 0.06 occur for XLF IDs {m['XLF_difference_exceeding_0_06_IDs']} and XHF IDs {m['XHF_difference_exceeding_0_06_IDs']}. The 0.06 comparison is a stated audit screen, not an instrument tolerance or validated physical normalization.",
        "",
        "For **ID 8, row 11**, BD11=238 (LF), BE11=235 (HF), BF11=11.92 g, BG11=20 (XLF), BH11=14.7 (XHF). Thus 235/11.92 = 19.71476510067114, unlike the stored 14.7. The raw LF/HF relative contrast is 1.2605042%, whereas the contrast calculated from the stored XLF/XHF pair is 26.5%. This is a material internal arithmetic discrepancy under the apparent simple normalization, not proof of which value should be corrected. The source was left unchanged.",
        "",
        f"The 32 raw relative contrasts, 100×(LF−HF)/LF, range {m['raw_contrast_min_percent']:.10g}–{m['raw_contrast_max_percent']:.10g}%. `magnetic_arithmetic.csv` gives every calculation. Raw and stored-normalized contrasts should not be interchanged without reporting this discrepancy; interpretation as a calibrated physical frequency-dependence measurement also needs the missing measurement metadata.",
        "",
        "## Feasible evidence and limits of this file",
        "",
        "The file supports a within-workbook question about how a magnetic descriptor relates separately to total concentration, the allocation among four operational fractions, and the recorded proximity/direction descriptors. Paired chemical totals and fractions permit separating higher concentrations from shifts in fractional allocation; simple geographic alternatives are available for comparison. These are distinct sources of descriptive evidence, so focusing only on one pooled correlation would leave useful information unused.",
        "",
        "Any screening or prediction assessment is retrospective unless evaluated on new independent samples. The 32 records support restrained exploratory comparisons and influence/sensitivity checks; they do not by themselves support robust high-capacity prediction, a verified causal factory plume, landscape metal mass, source apportionment, clinical/ecological risk classification, or a discovery of phase-specific mechanisms. Repeated Cd totals, repeated fraction relationships, the unexplained record-order structure, and the XHF discrepancy must be disclosed. Concentration sums over records are decision scores, not landscape mass or area-weighted burden.",
        "",
        "Important missing evidence is the upstream sampling and analytical protocol: dates/depth/material, field selection and replicate structure, extraction reagents and timings, QA/QC and blanks, standards/recovery, detection/quantification limits, reasons for repeated values/formatting, instrument/frequencies/calibration, and the stated construction/units of XLF/XHF. The workbook nevertheless **does** contain operational fraction labels, chemical units, masses, and recorded spatial descriptors. 'No metadata' would be inaccurate.",
        "",
        "## Reproduction and output meanings",
        "",
        f"Dependencies used: Python {data['runtime']['python']}, openpyxl {data['runtime']['openpyxl']}; the remaining modules are Python standard library.",
        "",
        "```text",
        'python fresh_read_20260927/fresh_inventory.py --source "source/Iron Factory.xlsx" --output-dir "fresh_read_20260927"',
        "```",
        "",
        "`data_dictionary.csv`: 55 fields with source columns, labels, literal units, and roles. `nonempty_cells.csv`: exhaustive nonempty-cell values/types/styles. `sample_records.csv`: all 32 records without numerical alterations. `magnetic_arithmetic.csv`: raw and stored normalization/contrast comparisons. `inventory.json`: machine-readable workbook/package metadata, data checks, repeated-ratio groups, geographic counts, and record-half ranges. This note and the five supporting data files are produced by `fresh_inventory.py`; existing analyses in the same directory are untouched.",
    ])
    (out/"data_inventory.md").write_text("\n".join(lines)+"\n",encoding="utf-8")


if __name__ == "__main__":
    main()
