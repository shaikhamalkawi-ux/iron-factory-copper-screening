"""Exploratory, all-element fresh-read diagnostics. Never writes the source workbook."""
from pathlib import Path
import argparse,hashlib,json,re,itertools,warnings
import numpy as np
import pandas as pd
import openpyxl
from scipy.stats import spearmanr
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

B=Path(__file__).resolve().parent
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--source',type=Path,default=Path(__file__).resolve().parents[1] / "source/Iron Factory.xlsx",help='Original workbook or unchanged copy (default: packaged source workbook).')
SOURCE=parser.parse_args().source
EXPECTED='1f403f49d399643455cffaad47d4445dee8acfba40ca14b6452283bd9e7e6b8b'
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==EXPECTED
s=openpyxl.load_workbook(SOURCE,data_only=True)['All rsults']
elements=['Cd','Co','Cr','Cu','Fe','Mn','Ni','Pb','Zn'];ids=np.arange(1,33)
F=np.array([[[s.cell(r,c+j).value for c in [2,12,22,32]] for j in range(9)] for r in range(4,36)],float)
T=F.sum(axis=2);P=F/T[:,:,None]
assert F.shape==(32,9,4) and np.all(F>0)
x=np.array([s.cell(r,59).value for r in range(4,36)],float)
dist=np.array([s.cell(r,53).value for r in range(4,36)],float)
LF,HF,mass,XHF=[np.array([s.cell(r,c).value for r in range(4,36)],float) for c in [56,57,58,60]]
directions=[str(s.cell(r,52).value).strip() for r in range(4,36)]
def dms(text):
    v=re.findall(r'\d+(?:\.\d+)?',str(text));assert len(v)==3
    return float(v[0])+float(v[1])/60+float(v[2])/3600
lat=np.array([dms(s.cell(r,54).value) for r in range(4,36)])
lon=np.array([dms(s.cell(r,55).value) for r in range(4,36)])
# Display-only local tangent approximation; origin is an analyst-chosen centroid.
lat0=lat[dist==0].mean();lon0=lon[dist==0].mean();R=6371008.8
east=R*np.cos(np.radians(lat0))*np.radians(lon-lon0);north=R*np.radians(lat-lat0)
radial=np.hypot(east,north)
meta=pd.DataFrame({'ID':ids,'direction':directions,'recorded_distance_m':dist,'latitude_deg':lat,'longitude_deg':lon,'east_relative_m':east,'north_relative_m':north,'centroid_distance_m':radial,'LF':LF,'HF':HF,'mass_g':mass,'stored_XLF':x,'stored_XHF':XHF,'LF_div_mass':LF/mass,'HF_div_mass':HF/mass,'raw_LF_HF_contrast_percent':100*(LF-HF)/LF,'stored_XLF_XHF_contrast_percent':100*(x-XHF)/x})
meta.to_csv(B/'metadata_diagnostics.csv',index=False,encoding='utf-8-sig')
def rho(a,b):return float(spearmanr(a,b).statistic)
def coef(y,mask,adjust=False):
    columns=[np.ones(sum(mask)),np.log(x[mask])]
    if adjust:columns.append(np.log1p(dist[mask]))
    design=np.column_stack(columns)
    assert np.linalg.matrix_rank(design)==design.shape[1]
    return float(np.linalg.lstsq(design,np.log(y[mask]),rcond=None)[0][1])

rows=[];identity_errors=[]
scenarios=[('all',np.ones(32,bool),False),('distance_adjusted',np.ones(32,bool),True),('outside_only',dist>0,False)]
for e,el in enumerate(elements):
    for scen,mask,adjust in scenarios:
        bt=coef(T[:,e],mask,adjust)
        for f in range(4):
            ba=coef(F[:,e,f],mask,adjust);bp=coef(P[:,e,f],mask,adjust)
            loo=[]
            for omitted in np.where(mask)[0]:
                reduced=mask.copy();reduced[omitted]=False
                loo.append((omitted+1,coef(F[:,e,f],reduced,adjust),coef(P[:,e,f],reduced,adjust)))
            identity_errors.append(abs(ba-bt-bp))
            rows.append({'element':el,'fraction':f'F{f+1}','scenario':scen,'n':int(mask.sum()),'beta_absolute':ba,'beta_total':bt,'beta_allocation':bp,'decomposition_error':ba-bt-bp,'rho_XLF_absolute':rho(x[mask],F[mask,e,f]),'rho_XLF_share':rho(x[mask],P[mask,e,f]),'loo_absolute_min':min(a[1] for a in loo),'loo_absolute_max':max(a[1] for a in loo),'loo_allocation_min':min(a[2] for a in loo),'loo_allocation_max':max(a[2] for a in loo)})
slopes=pd.DataFrame(rows);slopes.to_csv(B/'total_allocation_decomposition.csv',index=False)
assert max(identity_errors)<1e-12
summary=[]
for e,el in enumerate(elements):
    # Additive log-ratio against residual and an F1-versus-other-parts balance.
    balance=np.log(F[:,e,0])-np.log(F[:,e,1:]).mean(axis=1)
    out={'element':el,'total_min':T[:,e].min(),'total_median':np.median(T[:,e]),'total_max':T[:,e].max(),'rho_XLF_total':rho(x,T[:,e]),'rho_distance_total':rho(dist,T[:,e]),'rho_XLF_F1_vs_rest_balance':rho(x,balance),'fraction_share_unique_vectors_10decimals':len(np.unique(np.round(P[:,e,:],10),axis=0))}
    for f in range(4):out.update({f'F{f+1}_share_min':P[:,e,f].min(),f'F{f+1}_share_median':np.median(P[:,e,f]),f'F{f+1}_share_max':P[:,e,f].max()})
    for f in range(3):out[f'rho_XLF_log_F{f+1}_over_F4']=rho(x,np.log(F[:,e,f]/F[:,e,3]))
    summary.append(out)
pd.DataFrame(summary).to_csv(B/'element_summary.csv',index=False)

k=8;mag=np.lexsort((ids,-x))[:k]
cutoff=np.sort(dist)[k-1];fixed=np.where(dist<cutoff)[0];tied=np.where(dist==cutoff)[0]
nearest=[np.concatenate((fixed,np.array(c))) for c in itertools.combinations(tied,k-len(fixed))]
assert len(nearest)==6 and all(len(np.unique(a))==8 for a in nearest)
comparison=[];every_near=[]
for e,el in enumerate(elements):
    cf=F[:,e,:];oracle=np.sort(cf,axis=0)[-k:,:].sum(axis=0)
    outcomes={'total':T[:,e],'F1':cf[:,0],'F1_plus_F2_plus_F3':cf[:,:3].sum(axis=1)}
    for target,y in outcomes.items():
        o=np.sort(y)[-k:].sum();magscore=y[mag].sum()/o;ns=np.array([y[a].sum()/o for a in nearest])
        comparison.append({'element':el,'target':target,'magnetic_retention':magscore,'nearest_min_retention':min(ns),'nearest_max_retention':max(ns),'uniform_expected_retention':(k/32)*y.sum()/o})
        for a,score in zip(nearest,ns):every_near.append({'element':el,'target':target,'IDs':','.join(str(i+1) for i in sorted(a)),'retention':score})
    magscore=np.min(cf[mag].sum(axis=0)/oracle);ns=np.array([np.min(cf[a].sum(axis=0)/oracle) for a in nearest])
    comparison.append({'element':el,'target':'minimum_four_fraction','magnetic_retention':magscore,'nearest_min_retention':min(ns),'nearest_max_retention':max(ns),'uniform_expected_retention':None})
    for a,score in zip(nearest,ns):every_near.append({'element':el,'target':'minimum_four_fraction','IDs':','.join(str(i+1) for i in sorted(a)),'retention':score})
pd.DataFrame(comparison).to_csv(B/'magnetic_vs_nearest_k8.csv',index=False)
pd.DataFrame(every_near).to_csv(B/'nearest_all_tie_choices.csv',index=False)

# No mapped interpolation or inferred factory boundary is drawn.
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10})
fig,ax=plt.subplots(figsize=(8.5,6.4))
im=ax.scatter(east,north,c=x,cmap='viridis',s=72,edgecolors='white',linewidths=.6)
ax.scatter(east[dist==0],north[dist==0],s=135,facecolors='none',edgecolors='#d16127',linewidths=1.4,label='Recorded distance = 0')
for i in range(32):ax.annotate(str(i+1),(east[i],north[i]),xytext=(4,4),textcoords='offset points',fontsize=8)
ax.set(xlabel='Approximate east displacement (m)',ylabel='Approximate north displacement (m)',title='Recorded locations and stored XLF')
ax.set_aspect('equal');ax.grid(alpha=.18);ax.legend(loc='upper left',fontsize=8)
cb=fig.colorbar(im,ax=ax,shrink=.7);cb.set_label('Stored XLF (units unverified)')
fig.text(.07,.02,'Origin: centroid of the five distance-zero records; not a verified emission source.',fontsize=8)
fig.tight_layout(rect=[0,.05,1,1]);fig.savefig(B/'recorded_locations.png',dpi=200);plt.close(fig)

fig,axs=plt.subplots(1,2,figsize=(11,6.3),sharey=True)
matrices=[slopes[slopes.scenario=='all'].pivot(index='element',columns='fraction',values=c).reindex(elements).to_numpy() for c in ['beta_absolute','beta_allocation']]
lim=max(np.max(np.abs(a)) for a in matrices)
for ax,mat,title in zip(axs,matrices,['Absolute fraction concentration','Fraction share of element total']):
    im=ax.imshow(mat,cmap='RdBu_r',vmin=-lim,vmax=lim,aspect='auto')
    ax.set_xticks(range(4),['F1','F2','F3','F4']);ax.set_yticks(range(9),elements);ax.set_title(title,fontsize=11)
    for r in range(9):
        for c in range(4):ax.text(c,r,f'{mat[r,c]:.2f}',ha='center',va='center',fontsize=8,color='white' if abs(mat[r,c])>.60*lim else 'black')
fig.suptitle('Two distinct associations with stored XLF',fontsize=14)
fig.text(.08,.025,'Entries are exploratory log-log slopes, n=32. No causal or population inference; all 36 cases shown.',fontsize=9)
fig.subplots_adjust(left=.08,right=.87,wspace=.10,top=.86,bottom=.10)
fig.colorbar(im,cax=fig.add_axes([.9,.17,.02,.59]),label='Log-log slope')
fig.savefig(B/'absolute_vs_allocation_slopes.png',dpi=200);plt.close(fig)

v={'source_sha256':EXPECTED,'all_values_positive':bool(np.all(F>0)),'n_records':32,'n_elements':9,'n_fractions':4,'n_slope_rows':len(rows),'maximum_decomposition_error':max(identity_errors),'magnetic_IDs':(ids[mag]).tolist(),'nearest_tie_cutoff_m':float(cutoff),'nearest_fixed_IDs':ids[fixed].tolist(),'nearest_cutoff_tied_IDs':ids[tied].tolist(),'nearest_possible_sets':len(nearest),'rho_XLF_recorded_distance':rho(x,dist),'rho_XLF_recorded_distance_outside':rho(x[dist>0],dist[dist>0]),'rho_recorded_vs_centroid_distance':rho(dist,radial),'rho_XLF_LF_div_mass':rho(x,LF/mass),'same_top8_XLF_and_LF_div_mass':bool(set(mag)==set(np.lexsort((ids,-LF/mass))[:8])),'raw_LFHF_contrast_min_max_percent':[float(np.min(100*(LF-HF)/LF)),float(np.max(100*(LF-HF)/LF))],'source_unchanged':hashlib.sha256(SOURCE.read_bytes()).hexdigest()==EXPECTED,'no_pvalues_or_CIs_reported':True,'distances_are_recorded_covariates_not_validated_source_distances':True}
(B/'fresh_diagnostics_validation.json').write_text(json.dumps(v,indent=2),encoding='utf-8')
print(json.dumps(v,indent=2))
print('\nALL ELEMENT TOTALS');print(pd.DataFrame(summary)[['element','rho_XLF_total','rho_distance_total','rho_XLF_F1_vs_rest_balance','fraction_share_unique_vectors_10decimals']].to_string(index=False))
print('\nABSOLUTE / SHARE SIGN OPPOSITION: ALL36 TESTED, no significance filtering')
sel=slopes[(slopes.scenario=='all')&(slopes.beta_absolute*slopes.beta_allocation<0)]
print(sel[['element','fraction','beta_absolute','beta_total','beta_allocation','loo_allocation_min','loo_allocation_max']].to_string(index=False))
print('\nNEAREST COMPARISON MINIMUM FOUR FRACTIONS');print(pd.DataFrame(comparison).query("target=='minimum_four_fraction'").to_string(index=False))
