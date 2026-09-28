"""Follow-up sensitivity: all budgets and all boundary ties in both rules."""
from pathlib import Path
import argparse,itertools,json,hashlib
import numpy as np
import pandas as pd
import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
B=Path(__file__).resolve().parent
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--source',type=Path,default=Path(__file__).resolve().parents[1] / "source/Iron Factory.xlsx",help='Original workbook or unchanged copy (default: packaged source workbook).')
source=parser.parse_args().source
sha=hashlib.sha256(source.read_bytes()).hexdigest()
s=openpyxl.load_workbook(source,data_only=True)['All rsults']
F=np.array([[[s.cell(r,c+j).value for c in [2,12,22,32]] for j in range(9)] for r in range(4,36)],float)
elements=['Cd','Co','Cr','Cu','Fe','Mn','Ni','Pb','Zn']
x=np.array([s.cell(r,59).value for r in range(4,36)],float)
d=np.array([s.cell(r,53).value for r in range(4,36)],float)
def selections(v,k):
    cutoff=np.sort(v)[k-1];fixed=np.where(v<cutoff)[0];tied=np.where(v==cutoff)[0]
    return [np.r_[fixed,c].astype(int) for c in itertools.combinations(tied,k-len(fixed))]
rows=[];witnesses=[]
for k in range(1,32):
    mags=selections(-x,k);nears=selections(d,k)
    for j,el in enumerate(elements):
        oracle=np.sort(F[:,j,:],axis=0)[-k:].sum(axis=0)
        ms=np.array([F[a,j,:].sum(axis=0)/oracle for a in mags]);ns=np.array([F[a,j,:].sum(axis=0)/oracle for a in nears])
        mr=ms.min(axis=1);nr=ns.min(axis=1)
        classification='nearest_all_ties_better' if min(nr)>max(mr)+1e-12 else 'magnetic_all_ties_better' if min(mr)>max(nr)+1e-12 else 'overlap_or_equal'
        rows.append({'element':el,'k':k,'magnetic_choices':len(mags),'nearest_choices':len(nears),'magnetic_min':min(mr),'magnetic_max':max(mr),'nearest_min':min(nr),'nearest_max':max(nr),'classification':classification})
        if k==8:
            for a,n in zip(nears,ns):
                delta=n-ms[0]
                witnesses.append({'element':el,'nearest_IDs':','.join(str(i+1) for i in sorted(a)),'all_four_nearest_at_least_magnetic':bool(np.all(delta>=-1e-12)),'all_four_nearest_strictly_more':bool(np.all(delta>1e-12)),**{f'delta_F{f+1}':v for f,v in enumerate(delta)}})
df=pd.DataFrame(rows);df.to_csv(B/'proximity_vs_magnetic_all_budgets.csv',index=False)
pd.DataFrame(witnesses).to_csv(B/'k8_componentwise_rule_comparisons.csv',index=False)
tab=pd.crosstab(df.element,df.classification).reindex(elements,fill_value=0)
tab.to_csv(B/'proximity_budget_classification_counts.csv')
assert len(df)==279
# Check k=8 against separately generated first-pass results.
prior=pd.read_csv(B/'magnetic_vs_nearest_k8.csv').query("target=='minimum_four_fraction'").set_index('element')
for row in df.query('k==8').itertuples():
    p=prior.loc[row.element]
    assert np.isclose(row.magnetic_min,p.magnetic_retention,atol=1e-12,rtol=0)
    assert np.isclose(row.nearest_min,p.nearest_min_retention,atol=1e-12,rtol=0)
    assert np.isclose(row.nearest_max,p.nearest_max_retention,atol=1e-12,rtol=0)
assert hashlib.sha256(source.read_bytes()).hexdigest()==sha
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9})
fig,axs=plt.subplots(3,3,figsize=(11,8),sharex=True,sharey=True)
for ax,el in zip(axs.flat,elements):
    t=df[df.element==el]
    ax.fill_between(t.k,100*t.nearest_min,100*t.nearest_max,color='#d78b38',alpha=.22)
    ax.plot(t.k,100*t.nearest_min,color='#b46b1c',lw=1.1);ax.plot(t.k,100*t.nearest_max,color='#b46b1c',lw=1.1)
    ax.fill_between(t.k,100*t.magnetic_min,100*t.magnetic_max,color='#167f98',alpha=.22)
    ax.plot(t.k,100*t.magnetic_min,color='#167f98',lw=1.2);ax.plot(t.k,100*t.magnetic_max,color='#167f98',lw=1.2)
    ax.axvline(8,color='#666666',ls=':',lw=.8);ax.set_title(el);ax.set_ylim(0,101);ax.grid(alpha=.15)
for ax in axs[-1,:]:ax.set_xlabel('Selected records (k)')
for ax in axs[:,0]:ax.set_ylabel('Minimum retention (%)')
from matplotlib.lines import Line2D
fig.legend(handles=[Line2D([0],[0],color='#167f98',label='Highest XLF'),Line2D([0],[0],color='#b46b1c',label='Nearest recorded distance')],loc='upper center',ncol=2,bbox_to_anchor=(.5,.958),frameon=False)
fig.suptitle('Rule comparison changes with the budget',fontsize=15,y=.99)
fig.text(.07,.015,'Bands contain all cutoff-tie choices. They are not confidence intervals. Same-budget chemical oracles.',fontsize=9)
fig.tight_layout(rect=[0,.035,1,.92]);fig.savefig(B/'proximity_budget_profiles.png',dpi=200);plt.close(fig)
v={'source_unchanged':True,'rows':len(df),'all_nontrivial_budgets':31,'all_boundary_ties_enumerated_for_both_rules':True,'k8_agrees_with_first_script':True,'counts':json.loads(tab.to_json(orient='index'))}
(B/'proximity_budget_validation.json').write_text(json.dumps(v,indent=2),encoding='utf-8')
print(tab.to_string());print('\nK8 componentwise ordering counts')
print(pd.DataFrame(witnesses).groupby('element')[['all_four_nearest_at_least_magnetic','all_four_nearest_strictly_more']].sum().to_string())
