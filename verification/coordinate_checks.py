"""Conditional internal coordinate consistency; no physical source is assigned."""
from pathlib import Path
import re, json, itertools, hashlib, argparse
import numpy as np
from scipy.stats import pearsonr, spearmanr
import openpyxl
B=Path(__file__).resolve().parent
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--source',type=Path,default=Path(__file__).resolve().parents[1] / "source/Iron Factory.xlsx")
parser.add_argument('--output',type=Path,default=B)
args=parser.parse_args()
src=args.source
B=args.output
B.mkdir(parents=True,exist_ok=True)
assert hashlib.sha256(src.read_bytes()).hexdigest()=='1f403f49d399643455cffaad47d4445dee8acfba40ca14b6452283bd9e7e6b8b'
w=openpyxl.load_workbook(src,data_only=True)['All rsults']
def dms(s):
    nums=re.findall(r'\d+(?:\.\d+)?',s)
    assert len(nums)==3, s
    a,b,c=map(float,nums)
    assert 0<=b<60 and 0<=c<60
    return a+b/60+c/3600
coords=np.array([[dms(w.cell(r,j).value) for j in [54,55]] for r in range(4,36)])
ids=np.arange(1,33)
x=np.array([w.cell(r,59).value for r in range(4,36)])
d=np.array([w.cell(r,53).value for r in range(4,36)])
dirs=[w.cell(r,52).value.strip() for r in range(4,36)]
c=np.array([[w.cell(r,j).value for j in [5,15,25,35]] for r in range(4,36)])
targets=np.column_stack([c.sum(1),c[:,0],c[:,:3].sum(1)])
radians=np.radians(coords)
inside=np.flatnonzero(d==0)
def hav(a,b):
    da=b[0]-a[0]; db=b[1]-a[1]
    h=np.sin(da/2)**2+np.cos(a[0])*np.cos(b[0])*np.sin(db/2)**2
    return 6371008.8*2*np.arcsin(np.sqrt(np.clip(h,0,1)))
geo=np.array([min(hav(row,radians[j]) for j in inside) for row in radians])
positive=d>0
fit=np.polyfit(geo[positive],d[positive],1)
def choices(v,k,desc=False):
    cutoff=np.sort(v)[-k] if desc else np.sort(v)[k-1]
    fixed=np.flatnonzero(v>cutoff if desc else v<cutoff)
    ties=np.flatnonzero(v==cutoff)
    return [np.r_[fixed,np.array(t,dtype=int)] for t in itertools.combinations(ties,k-len(fixed))]
def comparison(k):
    oracle=np.sort(targets,axis=0)[-k:].sum(0)
    ms=np.array([targets[s].sum(0)/oracle for s in choices(x,k,True)])
    ps=np.array([targets[s].sum(0)/oracle for s in choices(geo,k)])
    margin=np.array([ms[:,0].min()-ps[:,0].max(),ms[:,1].min()-ps[:,1].max(),ps[:,2].min()-ms[:,2].max()])
    return {'k':k,'mag_min':ms.min(0).tolist(),'mag_max':ms.max(0).tolist(),'coordinate_min':ps.min(0).tolist(),'coordinate_max':ps.max(0).tolist(),'joint_reversal':bool(np.all(margin>1e-12))}
center=np.radians(coords[inside].mean(0))
bearings=[]
for lat,lon in radians:
    bearings.append((np.degrees(np.arctan2(np.sin(lon-center[1])*np.cos(lat),np.cos(center[0])*np.sin(lat)-np.sin(center[0])*np.cos(lat)*np.cos(lon-center[1])))+360)%360)
sectors=['شمال','شمال شرق','شرق','جنوب شرق','جنوب','جنوب غرب','غرب','شمال غرب']
pred=[sectors[int((b+22.5)//45)%8] for b in bearings]
exceptions=[{'id':int(ids[i]),'stored':dirs[i],'derived':pred[i],'bearing':bearings[i],'degrees_from_boundary':abs((bearings[i]+22.5)%45-0) if (bearings[i]+22.5)%45<22.5 else 45-(bearings[i]+22.5)%45} for i in range(32) if dirs[i]!=pred[i]]
out={'source_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),
     'assumptions':['Positive DMS latitude/longitude interpreted as decimal degrees; hemispheres and datum not certified','Spherical haversine with radius6371008.8m','Nearest of five zero-distance recorded points, not verified factory boundary or emission origin','Bearing sectors from arithmetic mean of five zero-distance coordinates, equal45degree sectors'],
     'positive_distance_count':int(positive.sum()),'pearson':float(pearsonr(d[positive],geo[positive]).statistic),'spearman':float(spearmanr(d[positive],geo[positive]).statistic),
     'mae_m':float(np.mean(np.abs(d[positive]-geo[positive]))),'rmse_m':float(np.sqrt(np.mean((d[positive]-geo[positive])**2))),
     'recorded_on_coordinate_slope':float(fit[0]),'intercept_m':float(fit[1]),'k8':comparison(8),'k8_coordinate_ids':[ids[s].tolist() for s in choices(geo,8)],
     'profiles':[comparison(k) for k in range(1,32)],'direction_agreements':32-len(exceptions),'direction_exceptions':exceptions,
     'coordinate_distance_by_id_m':geo.tolist()}
(B/'coordinate_checks.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:v for k,v in out.items() if k not in ['profiles','coordinate_distance_by_id_m']},ensure_ascii=False,indent=2))
