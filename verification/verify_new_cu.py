"""Read-only replay of the Cu checks from the original XLSX.

The workbook supplies every numerical computational input.
The table-only TeX fixture is read solely to compare its 81 printed entries.
All selection families are actually enumerated. A separate additive cutoff-bound
implementation checks their extrema. Output CSVs use ratios except fields ending
in _pct or _pp. Candidate omissions are conditional influence diagnostics.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import itertools
import json
import math
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET
import zipfile

import numpy as np
import openpyxl

EXPECTED_SHA = "1f403f49d399643455cffaad47d4445dee8acfba40ca14b6452283bd9e7e6b8b"
ELEMENTS = ["Cd", "Co", "Cr", "Cu", "Fe", "Mn", "Ni", "Pb", "Zn"]
TARGETS = ["total", "F1", "F123"]
TOL = 1e-12
CHECK_DIFFERENCES = []


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_csv(path, rows):
    assert rows, path
    with Path(path).open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def ids_text(indices):
    return ";".join(str(int(i) + 1) for i in sorted(indices))


def selections(rank, k):
    """Ascending rank: ties use exact stored numerical equality."""
    boundary = np.sort(rank)[k - 1]
    fixed = tuple(np.flatnonzero(rank < boundary))
    tied = tuple(np.flatnonzero(rank == boundary))
    h = k - len(fixed)
    family = [tuple(sorted(fixed + extra)) for extra in itertools.combinations(tied, h)]
    assert len(family) == math.comb(len(tied), h)
    assert all(len(s) == k and len(set(s)) == k for s in family)
    return family


def direct_bounds(rank, target, k):
    """Independent scalar implementation: additive tie extrema without enumeration."""
    cut = sorted(float(x) for x in rank)[k - 1]
    fixed = [float(y) for x, y in zip(rank, target) if x < cut]
    boundary = sorted(float(y) for x, y in zip(rank, target) if x == cut)
    h = k - len(fixed)
    oracle = math.fsum(sorted(map(float, target), reverse=True)[:k])
    low = (math.fsum(fixed) + math.fsum(boundary[:h])) / oracle
    high = (math.fsum(fixed) + math.fsum(boundary[-h:])) / oracle
    return low, high


def compare(x, distance, targets, k):
    """Return extrema, classifications and complete original-family local indices."""
    fm = selections(-x, k)
    fp = selections(distance, k)
    oracle = np.sort(targets, axis=0)[-k:].sum(axis=0)
    scores_m = np.array([targets[list(s)].sum(axis=0) / oracle for s in fm])
    scores_p = np.array([targets[list(s)].sum(axis=0) / oracle for s in fp])
    lm, um = scores_m.min(axis=0), scores_m.max(axis=0)
    lp, up = scores_p.min(axis=0), scores_p.max(axis=0)
    cls = np.where(lm > up + TOL, "M", np.where(lp > um + TOL, "P", "overlap"))
    for j in range(targets.shape[1]):
        bm = direct_bounds(-x, targets[:, j], k)
        bp = direct_bounds(distance, targets[:, j], k)
        diff = max(abs(a - b) for a, b in zip([lm[j], um[j], lp[j], up[j]], [*bm, *bp]))
        CHECK_DIFFERENCES.append(diff)
        assert diff < 1e-12, (k, j, diff)
    return {"oracle": oracle, "m_lo": lm, "m_hi": um, "p_lo": lp, "p_hi": up,
            "classification": cls, "family_m": fm, "family_p": fp}


def case_row(name, eligible, x, distance, targets, k, omitted=()):
    eligible = np.asarray(eligible, dtype=int)
    result = compare(x[eligible], distance[eligible], targets[eligible], k)
    row = {"case": name, "omitted_ids": ids_text(omitted), "eligible_ids": ids_text(eligible),
           "n": len(eligible), "k": k, "selected_fraction": k / len(eligible),
           "magnetic_choice_count": len(result["family_m"]),
           "proximity_choice_count": len(result["family_p"])}
    for j, target in enumerate(TARGETS):
        for field in ["oracle", "m_lo", "m_hi", "p_lo", "p_hi", "classification"]:
            val = result[field][j]
            row[f"{target}_{field}"] = float(val) if field != "classification" else str(val)
        row[f"{target}_M_worst_margin_pp"] = 100 * (result["m_lo"][j] - result["p_hi"][j])
        row[f"{target}_P_worst_margin_pp"] = 100 * (result["p_lo"][j] - result["m_hi"][j])
    row["three_target_pattern"] = list(result["classification"]) == ["M", "M", "P"]
    row["total_nonresidual_reversal"] = result["classification"][0] == "M" and result["classification"][2] == "P"
    return row


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=Path(__file__).resolve().parents[1] / "source/Iron Factory.xlsx")
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--baseline", type=Path, default=Path(__file__).resolve().parents[1] / "baseline/table2.tex")
    args = parser.parse_args()
    source_hash = sha(args.source)
    if source_hash != EXPECTED_SHA:
        raise ValueError(f"Source differs from preserved original: {source_hash}; expected {EXPECTED_SHA}")
    args.output.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(args.source, data_only=False, read_only=False)
    ws = wb["All rsults"]
    assert [ws.cell(r, 1).value for r in range(4, 36)] == list(range(1, 33))
    fractions = np.array([[[ws.cell(r, start + e).value for e in range(9)]
                          for start in [2, 12, 22, 32]] for r in range(4, 36)], dtype=float)
    reported = np.array([[ws.cell(r, 42 + e).value for e in range(9)] for r in range(4, 36)], dtype=float)
    magnetic = {name: np.array([ws.cell(r, col).value for r in range(4, 36)], dtype=float)
                for name, col in [("XLF", 59), ("LF", 56), ("HF", 57), ("mass", 58), ("XHF", 60)]}
    x = magnetic["XLF"]
    distance = np.array([ws.cell(r, 53).value for r in range(4, 36)], dtype=float)
    raw_directions = [ws.cell(r, 52).value for r in range(4, 36)]
    directions = np.array([d.strip() for d in raw_directions])
    assert np.isfinite(fractions).all() and (fractions > 0).all()
    assert np.isfinite(reported).all() and (reported > 0).all()
    assert (magnetic["mass"] > 0).all()
    all_targets = np.stack([fractions.sum(axis=1), fractions[:, 0, :], fractions[:, :3, :].sum(axis=1)], axis=2)
    cu = fractions[:, :, 3]
    targets = all_targets[:, 3, :]
    eligible_all = np.arange(32)

    # Independent cell-reading route through the XLSX XML numeric storage.
    with zipfile.ZipFile(args.source) as archive:
        xml = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    numeric_xml = {c.attrib["r"]: float(c.find("m:v", ns).text)
                   for c in xml.findall(".//m:sheetData/m:row/m:c", ns)
                   if c.find("m:v", ns) is not None and c.attrib.get("t", "n") == "n"}
    cell_rows = []
    for r in range(4, 36):
        for col in list(range(2, 11)) + list(range(12, 21)) + list(range(22, 31)) + list(range(32, 41)) + list(range(42, 51)) + [53, 56, 57, 58, 59, 60]:
            cell = ws.cell(r, col)
            assert cell.data_type != "f" and float(cell.value) == numeric_xml[cell.coordinate]
            cell_rows.append({"sheet": ws.title, "cell": cell.coordinate, "record_id": r - 3,
                              "value": cell.value, "xml_value": numeric_xml[cell.coordinate],
                              "source_sha256": source_hash})
    assert len(cell_rows) == 1632
    write_csv(args.output / "source_numeric_cells.csv", cell_rows)

    grid = []
    for k in range(1, 32):
        flat = all_targets.reshape(32, 27)
        result = compare(x, distance, flat, k)
        for e, element in enumerate(ELEMENTS):
            for q, target in enumerate(TARGETS):
                j = e * 3 + q
                row = {"element": element, "target": target, "k": k,
                       "magnetic_choice_count": len(result["family_m"]), "proximity_choice_count": len(result["family_p"])}
                for field in ["oracle", "m_lo", "m_hi", "p_lo", "p_hi", "classification"]:
                    row[field] = float(result[field][j]) if field != "classification" else str(result[field][j])
                grid.append(row)
    assert len(grid) == 837
    write_csv(args.output / "ordinary_target_grid.csv", grid)
    table = []
    tex = args.baseline.read_text(encoding="utf-8")
    for element in ELEMENTS:
        line = next(line for line in tex.splitlines() if line.startswith(element + " &"))
        printed = [float(x) for x in re.findall(r"\d+\.\d+", line)]
        expected = []
        for target in TARGETS:
            r = next(r for r in grid if r["element"] == element and r["target"] == target and r["k"] == 8)
            expected.extend([r["m_lo"] * 100, r["p_lo"] * 100, r["p_hi"] * 100])
        assert len(printed) == 9
        for j, (p, value) in enumerate(zip(printed, expected)):
            ok = f"{p:.2f}" == f"{value:.2f}"
            table.append({"element": element, "target": TARGETS[j // 3],
                          "statistic": ["magnetic", "proximity_min", "proximity_max"][j % 3],
                          "printed_pct": p, "recomputed_pct": value, "rounded_match": ok})
            assert ok, table[-1]
    write_csv(args.output / "r9_table2_replay.csv", table)

    singles = []
    for omitted in range(32):
        eligible = eligible_all[eligible_all != omitted]
        for k in range(1, 31):
            singles.append(case_row("single_omission", eligible, x, distance, targets, k, [omitted]))
    assert len(singles) == 960
    write_csv(args.output / "cu_single_omission_cases.csv", singles)
    single_counts = []
    for k in range(1, 31):
        rows = [r for r in singles if r["k"] == k]
        single_counts.append({"k": k, "three_target_count": sum(r["three_target_pattern"] for r in rows),
                              "two_endpoint_count": sum(r["total_nonresidual_reversal"] for r in rows),
                              "three_target_exception_ids": ";".join(r["omitted_ids"] for r in rows if not r["three_target_pattern"])})
    write_csv(args.output / "cu_single_omission_size_summary.csv", single_counts)
    singles8 = [r for r in singles if r["k"] == 8]
    pairs = []
    for omitted in itertools.combinations(range(32), 2):
        eligible = np.array([i for i in range(32) if i not in omitted])
        pairs.append(case_row("pair_omission", eligible, x, distance, targets, 8, omitted))
    assert len(pairs) == 496
    write_csv(args.output / "cu_pair_omission_cases.csv", pairs)
    pair_exceptions = [r for r in pairs if not r["three_target_pattern"]]
    write_csv(args.output / "cu_pair_omission_exceptions.csv", pair_exceptions)

    scope = [case_row("all_records", eligible_all, x, distance, targets, 8)]
    for k in [7, 8]:
        scope.append(case_row("positive_distance", eligible_all[distance > 0], x, distance, targets, k, eligible_all[distance == 0]))
    for name, eligible in [("IDs_1_16", np.arange(16)), ("IDs_17_32", np.arange(16, 32))]:
        for k in [4, 8]:
            scope.append(case_row(name, eligible, x, distance, targets, k, np.setdiff1d(eligible_all, eligible)))
    for d in sorted(set(directions)):
        scope.append(case_row("omit_direction_" + d, eligible_all[directions != d], x, distance, targets, 8, eligible_all[directions == d]))
    write_csv(args.output / "cu_candidate_scope_cases.csv", scope)

    proxy_rows = []
    base_family = selections(-x, 8)
    proxy_values = {"stored_XLF": x, "unrounded_LF_per_mass": magnetic["LF"] / magnetic["mass"],
                    "raw_LF": magnetic["LF"], "stored_XHF": magnetic["XHF"],
                    "raw_HF": magnetic["HF"], "unrounded_HF_per_mass": magnetic["HF"] / magnetic["mass"]}
    for name, proxy in proxy_values.items():
        family = selections(-proxy, 8)
        row = case_row(name, eligible_all, proxy, distance, targets, 8)
        row["magnetic_ids"] = "|".join(ids_text(s) for s in family)
        row["same_family_as_stored_XLF"] = set(family) == set(base_family)
        proxy_rows.append(row)
    write_csv(args.output / "cu_magnetic_proxy_k8.csv", proxy_rows)
    proxy_grid = [case_row(name, eligible_all, proxy, distance, targets, k)
                  for name, proxy in proxy_values.items() for k in range(1, 32)]
    write_csv(args.output / "cu_magnetic_proxy_all_sizes.csv", proxy_grid)
    six_names = ["F1", "F2", "F3", "F4", "F123", "total"]
    six_targets = np.column_stack((cu, cu[:, :3].sum(axis=1), cu.sum(axis=1)))
    six_grid = []
    for k in range(1, 32):
        result = compare(x, distance, six_targets, k)
        for j, name in enumerate(six_names):
            row = {"target": name, "k": k}
            for field in ["oracle", "m_lo", "m_hi", "p_lo", "p_hi", "classification"]:
                row[field] = float(result[field][j]) if field != "classification" else str(result[field][j])
            six_grid.append(row)
    write_csv(args.output / "cu_six_target_grid.csv", six_grid)
    write_csv(args.output / "cu_six_target_k8.csv", [r for r in six_grid if r["k"] == 8])

    # Six complete rows are retained so component extrema are never combined
    # across incompatible proximity choices.
    magnetic_set = set(base_family[0])
    near_family = selections(distance, 8)
    attribution = []
    contributions = []
    epsilon_rows = []
    identity_errors = []
    for choice, near_indices in enumerate(near_family, start=1):
        near_set = set(near_indices)
        shared = magnetic_set & near_set
        m_only, p_only = sorted(magnetic_set - near_set), sorted(near_set - magnetic_set)
        a = cu[m_only].sum(axis=0)
        b = cu[p_only].sum(axis=0)
        delta = a - b
        delta_nr = float(delta[:3].sum())
        delta_total = float(delta.sum())
        row = {"choice": choice, "proximity_ids": ids_text(near_set), "magnetic_ids": ids_text(magnetic_set),
               "shared_ids": ids_text(shared), "magnetic_only_ids": ids_text(m_only), "proximity_only_ids": ids_text(p_only)}
        for j in range(4):
            row[f"M_exclusive_F{j + 1}"] = float(a[j])
            row[f"P_exclusive_F{j + 1}"] = float(b[j])
            row[f"delta_F{j + 1}"] = float(delta[j])
        row.update(delta_F123=delta_nr, delta_total=delta_total,
                   lambda_star=-delta_nr / delta[3])
        assert delta[3] > 0 and 0 < row["lambda_star"] < 1
        # Compare exclusive and complete-selection calculations independently.
        full_delta = cu[sorted(magnetic_set)].sum(axis=0) - cu[sorted(near_set)].sum(axis=0)
        identity_errors.append(float(np.max(np.abs(full_delta - delta))))
        identity_errors.append(abs(delta_nr + delta[3] - delta_total))
        identity_errors.append(abs(delta_nr + row["lambda_star"] * delta[3]))
        attribution.append(row)
        for sign, side, records in [(1, "magnetic_only", m_only), (-1, "proximity_only", p_only)]:
            for i in records:
                record = {"choice": choice, "record_id": i + 1, "side": side, "sign": sign}
                for j in range(4):
                    record[f"signed_F{j + 1}"] = float(sign * cu[i, j])
                record["signed_F123"] = float(sign * cu[i, :3].sum())
                record["signed_total"] = float(sign * cu[i].sum())
                contributions.append(record)
        for q, target in enumerate(TARGETS):
            am = math.fsum(float(targets[i, q]) for i in m_only)
            bp = math.fsum(float(targets[i, q]) for i in p_only)
            eps = abs(am - bp) / (am + bp)
            winner, loser = max(am, bp), min(am, bp)
            equality_residual = (1 - eps) * winner - (1 + eps) * loser
            epsilon_rows.append({"choice": choice, "target": target, "M_exclusive_sum": am,
                                 "P_exclusive_sum": bp, "winner": "M" if am > bp else "P",
                                 "epsilon_star": eps, "epsilon_star_pct": 100 * eps,
                                 "equality_residual": equality_residual})
            assert abs(equality_residual) < 1e-10
    assert max(identity_errors) < 1e-10
    write_csv(args.output / "cu_k8_fraction_attribution.csv", attribution)
    write_csv(args.output / "cu_k8_signed_record_contributions.csv", contributions)
    write_csv(args.output / "cu_k8_hypothetical_epsilon.csv", epsilon_rows)

    # Continuation scores, recomputing the full chemical oracle for each lambda.
    lambda_grid = []
    lambda_values = sorted(set([j / 100 for j in range(101)] + [r["lambda_star"] for r in attribution]))
    for lam in lambda_values:
        target = cu[:, :3].sum(axis=1) + lam * cu[:, 3]
        result = compare(x, distance, target[:, None], 8)
        lambda_grid.append({"lambda": lam, "oracle": float(result["oracle"][0]),
                            "m_lo": float(result["m_lo"][0]), "m_hi": float(result["m_hi"][0]),
                            "p_lo": float(result["p_lo"][0]), "p_hi": float(result["p_hi"][0]),
                            "classification": str(result["classification"][0])})
    write_csv(args.output / "cu_k8_lambda_profile.csv", lambda_grid)

    minimal = {}
    for target, direction in [("total", "M"), ("F1", "M"), ("F123", "P")]:
        field = f"{target}_{direction}_worst_margin_pp"
        r = min(singles8, key=lambda r: r[field])
        minimal[target] = {"minimum_winning_margin_pp": r[field], "omitted_id": int(r["omitted_ids"])}
    counts = {"single_omission_size_cases": len(singles), "single_omission_k8_three_target": sum(r["three_target_pattern"] for r in singles8),
              "pair_omission_cases": len(pairs), "pair_omission_three_target": sum(r["three_target_pattern"] for r in pairs),
              "pair_omission_two_endpoint": sum(r["total_nonresidual_reversal"] for r in pairs),
              "pair_omission_F123_proximity": sum(r["F123_classification"] == "P" for r in pairs)}
    assert counts == {"single_omission_size_cases": 960, "single_omission_k8_three_target": 32,
                      "pair_omission_cases": 496, "pair_omission_three_target": 490,
                      "pair_omission_two_endpoint": 494, "pair_omission_F123_proximity": 496}
    assert {r["omitted_ids"] for r in pair_exceptions} == {"14;15", "14;26", "15;19", "15;24", "15;26", "15;29"}
    assert all(r["same_family_as_stored_XLF"] for r in proxy_rows)
    assert sha(args.source) == source_hash
    summary = {
        "source": {"path": str(args.source.resolve()), "sha256": source_hash, "hash_after_analysis": sha(args.source),
                   "sheet": ws.title, "record_rows": "4:35", "fraction_ranges": ["B4:J35", "L4:T35", "V4:AD35", "AF4:AN35"],
                   "total_range": "AP4:AX35", "Cu_fraction_columns": ["E", "O", "Y", "AI"], "Cu_reported_total_column": "AS",
                   "record_id_column": "A", "direction_column": "AZ", "distance_column": "BA", "LF_column": "BD", "HF_column": "BE",
                   "mass_column": "BF", "XLF_column": "BG", "XHF_column": "BH", "source_formulas_count": sum(c.data_type == "f" for row in ws for c in row),
                   "directions_raw": raw_directions, "directions_grouping": "Only strip surrounding whitespace for diagnostic groups; no source change.",
                   "LF_per_mass_max_abs_difference_from_XLF": float(np.max(np.abs(magnetic["LF"] / magnetic["mass"] - x))),
                   "record8": {name: float(val[7]) for name, val in magnetic.items()},
                   "record8_HF_per_mass": float(magnetic["HF"][7] / magnetic["mass"][7])},
        "method": {"comparison_tolerance_ratio": TOL, "enumeration": "Every admissible exact-value cutoff subset is enumerated for both rules.",
                   "ordinary_targets": TARGETS, "oracle": "Recomputed top-k target sum within each eligible candidate set.",
                   "fixed_k_qualification": "Omissions change selected proportion. Half checks explicitly include 4/16 and 8/16.",
                   "inference": "Conditional finite-table influence and candidate-scope checks, not external validation, sampling intervals, or future probabilities.",
                   "epsilon": "Independent relative target perturbations bounded by epsilon, fixed selections; shared records use the same perturbed values and cancel. Exact equality radius = abs(exclusive difference)/(exclusive sums). No measured error distribution or calibration is assumed.",
                   "lambda": "F123 + lambda*F4; preference sensitivity only. Each profile point recomputes the full-set chemical oracle.",
                   "not_replayed": "Old maximin optimization and exhaustive uniform-subset distributions were not rerun."},
        "verification": {"xml_openpyxl_numeric_cells_identical": len(cell_rows), "ordinary_grid_cases": len(grid), "printed_R9_table2_values_matching": len(table),
                         "max_total_reconstruction_difference": float(np.max(np.abs(reported - fractions.sum(axis=1)))),
                         "enumeration_vs_scalar_bound_checks": len(CHECK_DIFFERENCES), "max_enumeration_vs_scalar_difference_ratio": max(CHECK_DIFFERENCES),
                         "max_attribution_identity_difference": max(identity_errors), "source_unchanged": True,
                         "python": sys.version.split()[0], "numpy": np.__version__, "openpyxl": openpyxl.__version__},
        "counts": counts,
        "single_k8_minimum_margins": minimal,
        "sizes_preserving_three_targets_after_every_single_omission": [r["k"] for r in single_counts if r["three_target_count"] == 32],
        "single_size_counts": single_counts,
        "pair_exceptions": [{k: r[k] for k in ["omitted_ids", "total_classification", "F1_classification", "F123_classification", "total_M_worst_margin_pp", "total_P_worst_margin_pp", "F1_M_worst_margin_pp", "F1_P_worst_margin_pp", "F123_P_worst_margin_pp"]} for r in pair_exceptions],
        "candidate_scope": scope,
        "fraction_delta_ranges": {name: [min(r[name] for r in attribution), max(r[name] for r in attribution)] for name in ["delta_F1", "delta_F2", "delta_F3", "delta_F4", "delta_F123", "delta_total"]},
        "lambda_threshold_range": [min(r["lambda_star"] for r in attribution), max(r["lambda_star"] for r in attribution)],
        "hypothetical_epsilon_minima_pct": {q: min(r["epsilon_star_pct"] for r in epsilon_rows if r["target"] == q) for q in TARGETS},
        "all_k8_magnetic_proxy_families_identical": True,
        "magnetic_proxy_three_target_sizes": {name: [r["k"] for r in proxy_grid if r["case"] == name and r["three_target_pattern"]] for name in proxy_values},
        "cu_six_target_k8": [r for r in six_grid if r["k"] == 8],
        "cu_six_target_sizes_by_winner": {name: {winner: [r["k"] for r in six_grid if r["target"] == name and r["classification"] == winner] for winner in ["M", "P", "overlap"]} for name in six_names},
        "cu_fraction_shares_of_all_record_total_pct": {f"F{j + 1}": 100 * float(cu[:, j].sum() / cu.sum()) for j in range(4)},
        "raw_LF_HF_relative_difference_pct": {"min": float(np.min(100 * (magnetic["LF"] - magnetic["HF"]) / magnetic["LF"])),
                                              "max": float(np.max(100 * (magnetic["LF"] - magnetic["HF"]) / magnetic["LF"])),
                                              "mean": float(np.mean(100 * (magnetic["LF"] - magnetic["HF"]) / magnetic["LF"])),
                                              "median": float(np.median(100 * (magnetic["LF"] - magnetic["HF"]) / magnetic["LF"]))},
    }
    (args.output / "verification_summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8")
    manifest = {p.name: sha(p) for p in sorted(args.output.iterdir()) if p.is_file() and p.suffix in [".csv", ".json", ".py"] and p.name != "verification_manifest.json"}
    (args.output / "verification_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"counts": counts, "single_k8_minimum_margins": minimal, "pair_exceptions": summary["pair_exceptions"],
                      "lambda_threshold_range": summary["lambda_threshold_range"], "hypothetical_epsilon_minima_pct": summary["hypothetical_epsilon_minima_pct"],
                      "source_sha256": source_hash, "verification": summary["verification"], "output": str(args.output.resolve())}, indent=2))


if __name__ == "__main__":
    main()
